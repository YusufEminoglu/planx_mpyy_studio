"""Final-check workbook: every non-obvious SLD mapping with the decision already taken and why.

    py -3 tools/build_plangml_sld_review.py <output.xlsx>

Rules whose code-list position and title agree are the Ministry's own mapping
and are not listed. Everything else was decided in ``styles/plangml_sld/decisions.json``
and is listed here with its reason, so the author only marks an objection.
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SLD_DIR = ROOT / "styles" / "plangml_sld"


def main(output):
    mapping = json.loads((SLD_DIR / "mapping.json").read_text(encoding="utf-8"))
    decisions = json.loads((SLD_DIR / "decisions.json").read_text(encoding="utf-8"))
    null_reasons = {(r["level"], r["type"], r["literal"]): r["gerekce"] for r in decisions["literal_rules"]}
    rows = []
    for level, report in mapping["levels"].items():
        for type_name, entry in report["types"].items():
            if entry["style_match"] == "abbreviated":
                rows.append([level, type_name, entry["style"], "", "", "", "Stil bağlandı",
                             decisions["abbreviated_style_names"]["gerekce"]])
            elif entry["style_match"] == "decision":
                alias = decisions["style_aliases"][level][type_name]
                rows.append([level, type_name, entry["style"], "", "", "", "Yakın resmî stil bağlandı", alias["gerekce"]])
            for rule in entry["rules"]:
                for c in rule["conditions"]:
                    resolution = c.get("resolution")
                    if resolution == "title":
                        rows.append([level, type_name, entry["style"], rule["title"], c["by_index"], c["code"],
                                     "Başlıktaki kod uygulandı", decisions["title_over_index"]["gerekce"]])
                    elif resolution == "unresolved":
                        rows.append([level, type_name, entry["style"], rule["title"], "(kod listesinde yok)", "",
                                     "Kural atlandı", decisions["unresolved_rules"]["gerekce"]])
                    elif resolution == "decision_is_null":
                        rows.append([level, type_name, entry["style"], rule["title"], c["literal"], f'{c["field"]} boş',
                                     "Boş değer kuralı", null_reasons[(level, type_name, c["literal"])]])
        for type_name in report["missing_styles"]:
            symbol = decisions["catalog_symbols"].get(level, {}).get(type_name)
            if symbol:
                rows.append([level, type_name, "(detay kataloğu)", symbol["tanim"], "", "", "Katalog tanımından sembol",
                             symbol["gerekce"] + " Kaynak: " + symbol["kaynak"]])
            else:
                rows.append([level, type_name, "", "", "", "", "Nötr kaldı", "Ne SLD ne detay kataloğu tanımı bulundu."])
    for level in ("UIP", "NIP"):
        rows.append([level, "Yolorta", f"{level}_YOLORTA", "", "", "", "Nötr kaldı",
                     "SLD yalnız yol genişliği yazısı içeriyor; QGIS bundan çizim kuralı kuramıyor. Yol orta çizgisi gri çizilir."])

    book = Workbook()
    guide = book.active
    guide.title = "Açıklama"
    for line in [
        "PlanGML → resmî e-Plan gösterimi: son kontrol",
        "",
        "Kararlar verildi ve uygulandı. Yalnız katılmadığınız satırda 'İtiraz' sütununu seçin, isterseniz 'Not' yazın.",
        "Listelenmeyen kurallar Bakanlığın kendi eşleşmesidir (sıra numarası ile kural başlığı aynı kodu gösteriyor).",
        "",
        f"Listelenen karar: {len(rows)}",
        "Kaynaklar: PlanGML 1.1.7 XSD (eplan.csb.gov.tr), 'E-Plan SLD.zip' GeoServer stil seti, MPYY detay kataloğu.",
    ]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 120

    sheet = book.create_sheet("Son kontrol")
    sheet.append(["Kademe", "PlanGML tipi", "SLD stili / kaynak", "Kural başlığı / tanım", "Sıra numarasının kodu",
                  "Uygulanan", "Karar", "Gerekçe", "İtiraz", "Not"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        sheet.append(row + ["", ""])
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = sheet.max_row
    objection = DataValidation(type="list", formula1='"Doğru değil,Emin değilim"', allow_blank=True)
    sheet.add_data_validation(objection)
    objection.add(f"I2:I{last}")
    sheet.conditional_formatting.add(f"A2:J{last}", FormulaRule(formula=["LEN($I2)>0"], fill=PatternFill("solid", fgColor="F8CBAD")))
    for index, width in enumerate([8, 30, 28, 34, 26, 26, 24, 70, 16, 30], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:J{last}"

    absent = book.create_sheet("Eksik tarama görselleri")
    absent.append(["Dosya", "Not"])
    for name in mapping["images"]["absent_from_source"]:
        absent.append([name, "SLD'de atıf var, Bakanlık arşivinde dosya yok; ilgili dolgu boş çizilir"])
    absent.column_dimensions["A"].width = 44
    absent.column_dimensions["B"].width = 70
    book.save(output)
    print("rows", len(rows))


if __name__ == "__main__":
    main(sys.argv[1])
